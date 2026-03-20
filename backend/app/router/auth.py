from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from sqlalchemy.orm import Session
from app.database import get_db
from app.models.user import User
from app.schema.user import (
    UserCreate,
    UserUpdate,
    PasswordChange,
    UserLogin,
    UserResponse,
    UserListResponse,
    Token,
    TokenRefresh,
    ImportResult,
)
from app.service.auth import AuthService
from openpyxl import load_workbook
import io

router = APIRouter()
auth_service = AuthService()


@router.post("/register", response_model=UserResponse)
def register(user: UserCreate, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.username == user.username).first()
    if db_user:
        raise HTTPException(status_code=400, detail="Username already registered")

    if user.email:
        db_user = db.query(User).filter(User.email == user.email).first()
        if db_user:
            raise HTTPException(status_code=400, detail="Email already registered")

    if user.phone:
        db_user = db.query(User).filter(User.phone == user.phone).first()
        if db_user:
            raise HTTPException(status_code=400, detail="Phone already registered")

    hashed_password = auth_service.get_password_hash(user.password)
    db_user = User(
        username=user.username,
        email=user.email,
        phone=user.phone,
        hashed_password=hashed_password,
    )
    db.add(db_user)
    db.commit()
    db.refresh(db_user)
    return db_user


@router.post("/login", response_model=Token)
def login(user: UserLogin, db: Session = Depends(get_db)):
    db_user = (
        db.query(User).filter(User.username == user.account).first()
        or db.query(User).filter(User.email == user.account).first()
        or db.query(User).filter(User.phone == user.account).first()
    )

    if not db_user or not auth_service.authenticate_user(db_user, user.password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Incorrect account or password",
            headers={"WWW-Authenticate": "Bearer"},
        )
    access_token = auth_service.create_access_token(data={"sub": db_user.username})
    refresh_token = auth_service.create_refresh_token(data={"sub": db_user.username})
    return {
        "access_token": access_token,
        "refresh_token": refresh_token,
        "token_type": "bearer",
    }


@router.post("/refresh", response_model=Token)
def refresh(token_data: TokenRefresh, db: Session = Depends(get_db)):
    payload = auth_service.verify_refresh_token(token_data.refresh_token)
    if payload is None:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid or expired refresh token",
            headers={"WWW-Authenticate": "Bearer"},
        )
    username = payload.get("sub")
    if username is None:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid refresh token",
        )
    db_user = db.query(User).filter(User.username == username).first()
    if not db_user:
        raise HTTPException(status_code=401, detail="User not found")
    access_token = auth_service.create_access_token(data={"sub": db_user.username})
    refresh_token = auth_service.create_refresh_token(data={"sub": db_user.username})
    return {
        "access_token": access_token,
        "refresh_token": refresh_token,
        "token_type": "bearer",
    }


@router.get("/users", response_model=UserListResponse)
def get_users(
    skip: int = 0, limit: int = 100, keyword: str = None, db: Session = Depends(get_db)
):
    query = db.query(User)
    if keyword:
        keyword = keyword.strip()
        if keyword:
            query = query.filter(
                (User.username.contains(keyword))
                | (User.email.contains(keyword))
                | (User.phone.contains(keyword))
            )
    total = query.count()
    users = query.offset(skip).limit(limit).all()
    return {"total": total, "users": users}


@router.get("/users/{user_id}", response_model=UserResponse)
def get_user(user_id: int, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")
    return db_user


@router.put("/users/{user_id}", response_model=UserResponse)
def update_user(user_id: int, user_update: UserUpdate, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")

    if user_update.username and user_update.username != db_user.username:
        existing = db.query(User).filter(User.username == user_update.username).first()
        if existing:
            raise HTTPException(status_code=400, detail="Username already exists")
        db_user.username = user_update.username

    if user_update.email is not None and user_update.email != db_user.email:
        existing = db.query(User).filter(User.email == user_update.email).first()
        if existing:
            raise HTTPException(status_code=400, detail="Email already exists")
        db_user.email = user_update.email

    if user_update.phone is not None and user_update.phone != db_user.phone:
        existing = db.query(User).filter(User.phone == user_update.phone).first()
        if existing:
            raise HTTPException(status_code=400, detail="Phone already exists")
        db_user.phone = user_update.phone

    if user_update.is_active is not None:
        db_user.is_active = user_update.is_active

    db.commit()
    db.refresh(db_user)
    return db_user


@router.post("/users/{user_id}/change-password")
def change_password(
    user_id: int, password_data: PasswordChange, db: Session = Depends(get_db)
):
    db_user = db.query(User).filter(User.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")

    db_user.hashed_password = auth_service.get_password_hash(password_data.new_password)
    db.commit()
    return {"message": "Password updated successfully"}


@router.delete("/users/{user_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_user(user_id: int, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")
    db.delete(db_user)
    db.commit()


DEFAULT_PASSWORD = "123456"


@router.post("/users/import", response_model=ImportResult)
async def import_users(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not (file.filename.endswith(".xlsx") or file.filename.endswith(".xls")):
        raise HTTPException(
            status_code=400, detail="Only Excel files (.xlsx, .xls) are supported"
        )

    content = await file.read()
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    if ws is None:
        raise HTTPException(status_code=400, detail="Excel file has no data")

    headers = [cell.value for cell in ws[1]]
    username_idx = None
    email_idx = None
    phone_idx = None

    for i, h in enumerate(headers):
        if h == "username":
            username_idx = i
        elif h == "email":
            email_idx = i
        elif h == "phone":
            phone_idx = i

    if username_idx is None:
        raise HTTPException(status_code=400, detail="Excel must have 'username' column")

    total = 0
    success = 0
    failed = 0
    errors = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue

        total += 1
        username = (
            str(row[username_idx]).strip() if row[username_idx] is not None else ""
        )
        email = (
            str(row[email_idx]).strip()
            if email_idx is not None and row[email_idx] is not None
            else ""
        )
        phone = (
            str(row[phone_idx]).strip()
            if phone_idx is not None and row[phone_idx] is not None
            else ""
        )

        email = email if email else None
        phone = phone if phone else None

        if not username:
            failed += 1
            errors.append(f"Row {total + 1}: Username is required")
            continue

        existing = db.query(User).filter(User.username == username).first()
        if existing:
            failed += 1
            errors.append(f"Row {total + 1}: Username '{username}' already exists")
            continue

        if email:
            existing = db.query(User).filter(User.email == email).first()
            if existing:
                failed += 1
                errors.append(f"Row {total + 1}: Email '{email}' already exists")
                continue

        if phone:
            existing = db.query(User).filter(User.phone == phone).first()
            if existing:
                failed += 1
                errors.append(f"Row {total + 1}: Phone '{phone}' already exists")
                continue

        hashed_password = auth_service.get_password_hash(DEFAULT_PASSWORD)
        new_user = User(
            username=username,
            email=email,
            phone=phone,
            hashed_password=hashed_password,
        )
        db.add(new_user)
        success += 1

    db.commit()
    return ImportResult(total=total, success=success, failed=failed, errors=errors)
