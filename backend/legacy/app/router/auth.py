import io
import logging
from typing import Dict, Optional

from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from sqlalchemy.orm import Session
from app.database import get_db
from app.models.role import Role
from app.models.user_role import UserRole
from app.models.user import User
from app.schema.user import (
    UserCreate,
    UserUpdate,
    PasswordChange,
    UserLogin,
    UserResponse,
    UserListResponse,
    Token,
    LoginResponse,
    TokenRefresh,
    ImportResult,
)
from app.service.auth import AuthService
from app.service.router import RouterService
from openpyxl import load_workbook

logger = logging.getLogger("uvicorn.error")

router = APIRouter()
auth_service = AuthService()
router_service = RouterService()


def _validate_role_id(db: Session, role_id: Optional[str]) -> Optional[str]:
    if not role_id:
        return None
    existing = db.query(Role).filter(Role.id == role_id).first()
    if not existing:
        raise HTTPException(status_code=400, detail="Role does not exist")
    return role_id


def _replace_user_role(db: Session, user_id: str, role_id: Optional[str]) -> None:
    db.query(UserRole).filter(UserRole.user_id == user_id).delete()
    if role_id:
        db.add(UserRole(user_id=user_id, role_id=role_id))


def _build_user_role_map(db: Session, user_ids: list[str]) -> Dict[str, Optional[Dict[str, str]]]:
    if not user_ids:
        return {}

    rows = (
        db.query(UserRole.user_id, Role.id, Role.name)
        .join(Role, Role.id == UserRole.role_id)
        .filter(UserRole.user_id.in_(user_ids))
        .all()
    )
    result: Dict[str, Optional[Dict[str, str]]] = {user_id: None for user_id in user_ids}
    for user_id, role_id, role_name in rows:
        result[user_id] = {"id": role_id, "name": role_name}
    return result


def _serialize_user(user: User, role: Optional[Dict[str, str]]) -> Dict:
    return {
        "id": user.id,
        "username": user.username,
        "email": user.email,
        "phone": user.phone,
        "is_active": user.is_active,
        "is_superuser": user.is_superuser,
        "role": role,
        "role_id": role["id"] if role else None,
    }


def _collect_user_route_keys(db: Session, user_id: str):
    rows = (
        db.query(Role.route_keys)
        .join(UserRole, UserRole.role_id == Role.id)
        .filter(UserRole.user_id == user_id)
        .all()
    )
    for (keys,) in rows:
        if isinstance(keys, list):
            return [str(key) for key in keys if key]
    return []


@router.post("/register", response_model=UserResponse)
def register(user: UserCreate, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.username == user.username).first()
    if db_user:
        raise HTTPException(status_code=400, detail="Username already registered")

    if user.email:
        db_user = db.query(User).filter(User.email == user.email).first()
        if db_user:
            raise HTTPException(status_code=400, detail="Email already registered")

    if user.phone:
        db_user = db.query(User).filter(User.phone == user.phone).first()
        if db_user:
            raise HTTPException(status_code=400, detail="Phone already registered")

    valid_role_id = _validate_role_id(db, user.role_id)

    hashed_password = auth_service.get_password_hash(user.password)
    db_user = User(
        username=user.username,
        email=user.email,
        phone=user.phone,
        hashed_password=hashed_password,
    )
    db.add(db_user)
    db.commit()
    _replace_user_role(db, db_user.id, valid_role_id)
    db.commit()
    db.refresh(db_user)
    role_map = _build_user_role_map(db, [db_user.id])
    return _serialize_user(db_user, role_map.get(db_user.id))


@router.post("/login", response_model=LoginResponse)
def login(user: UserLogin, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.email == user.email).first()

    if not db_user or not auth_service.authenticate_user(db_user, user.password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Incorrect email or password",
            headers={"WWW-Authenticate": "Bearer"},
        )
    access_token = auth_service.create_access_token(data={"sub": db_user.username})
    refresh_token = auth_service.create_refresh_token(data={"sub": db_user.username})
    logger.info("User %s logged in successfully %s", db_user.username, db_user.is_superuser)
    if db_user.username == "admin":
        db_user.is_superuser = True
        db.commit()
        db.refresh(db_user)
    role = "admin" if db_user.is_superuser else "client"
    role_map = _build_user_role_map(db, [db_user.id])
    user_role = role_map.get(db_user.id)
    user_route_keys = _collect_user_route_keys(db, db_user.id)
    if not db_user.is_superuser and not user_route_keys:
        # Backward compatible fallback for users not yet bound to any role.
        user_route_keys = ["ClientHome", "ClientContents", "ClientMyInstances"]

    return {
        "access_token": access_token,
        "refresh_token": refresh_token,
        "token_type": "bearer",
        "user": {
            "id": db_user.id,
            "username": db_user.username,
            "email": db_user.email,
            "phone": db_user.phone,
            "is_active": db_user.is_active,
            "is_superuser": db_user.is_superuser,
            "role": role,
            "user_role": user_role,
            "role_id": user_role["id"] if user_role else None,
        },
        "routes": router_service.get_routes(db_user, user_route_keys),
        "menu": router_service.get_menu(db_user, user_route_keys),
    }


@router.post("/refresh", response_model=Token)
def refresh(token_data: TokenRefresh, db: Session = Depends(get_db)):
    payload = auth_service.verify_refresh_token(token_data.refresh_token)
    if payload is None:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid or expired refresh token",
            headers={"WWW-Authenticate": "Bearer"},
        )
    username = payload.get("sub")
    if username is None:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid refresh token",
        )
    db_user = db.query(User).filter(User.username == username).first()
    if not db_user:
        raise HTTPException(status_code=401, detail="User not found")
    access_token = auth_service.create_access_token(data={"sub": db_user.username})
    refresh_token = auth_service.create_refresh_token(data={"sub": db_user.username})
    return {
        "access_token": access_token,
        "refresh_token": refresh_token,
        "token_type": "bearer",
    }


@router.get("/users", response_model=UserListResponse)
def get_users(
    skip: int = 0, limit: int = 100, keyword: str = None, db: Session = Depends(get_db)
):
    query = db.query(User)
    if keyword:
        keyword = keyword.strip()
        if keyword:
            query = query.filter(
                (User.username.contains(keyword))
                | (User.email.contains(keyword))
                | (User.phone.contains(keyword))
            )
    total = query.count()
    users = query.offset(skip).limit(limit).all()
    role_map = _build_user_role_map(db, [item.id for item in users])
    serialized_users = [_serialize_user(item, role_map.get(item.id)) for item in users]
    return {"total": total, "users": serialized_users}


@router.get("/users/{user_id}", response_model=UserResponse)
def get_user(user_id: str, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")
    role_map = _build_user_role_map(db, [db_user.id])
    return _serialize_user(db_user, role_map.get(db_user.id))


@router.put("/users/{user_id}", response_model=UserResponse)
def update_user(user_id: str, user_update: UserUpdate, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")

    valid_role_id = None
    should_update_role = "role_id" in user_update.model_fields_set
    if should_update_role:
        valid_role_id = _validate_role_id(db, user_update.role_id)

    if user_update.username and user_update.username != db_user.username:
        existing = db.query(User).filter(User.username == user_update.username).first()
        if existing:
            raise HTTPException(status_code=400, detail="Username already exists")
        db_user.username = user_update.username

    if user_update.email is not None and user_update.email != db_user.email:
        existing = db.query(User).filter(User.email == user_update.email).first()
        if existing:
            raise HTTPException(status_code=400, detail="Email already exists")
        db_user.email = user_update.email

    if user_update.phone is not None and user_update.phone != db_user.phone:
        existing = db.query(User).filter(User.phone == user_update.phone).first()
        if existing:
            raise HTTPException(status_code=400, detail="Phone already exists")
        db_user.phone = user_update.phone

    if user_update.is_active is not None:
        db_user.is_active = user_update.is_active

    db.commit()
    if should_update_role:
        _replace_user_role(db, db_user.id, valid_role_id)
        db.commit()
    db.refresh(db_user)
    role_map = _build_user_role_map(db, [db_user.id])
    return _serialize_user(db_user, role_map.get(db_user.id))


@router.post("/users/{user_id}/change-password")
def change_password(
    user_id: str, password_data: PasswordChange, db: Session = Depends(get_db)
):
    db_user = db.query(User).filter(User.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")

    db_user.hashed_password = auth_service.get_password_hash(password_data.new_password)
    db.commit()
    return {"message": "Password updated successfully"}


@router.delete("/users/{user_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_user(user_id: str, db: Session = Depends(get_db)):
    db_user = db.query(User).filter(User.id == user_id).first()
    if not db_user:
        raise HTTPException(status_code=404, detail="User not found")
    db.query(UserRole).filter(UserRole.user_id == user_id).delete()
    db.delete(db_user)
    db.commit()


DEFAULT_PASSWORD = "123456"


@router.post("/users/import", response_model=ImportResult)
async def import_users(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if not (file.filename.endswith(".xlsx") or file.filename.endswith(".xls")):
        raise HTTPException(
            status_code=400, detail="Only Excel files (.xlsx, .xls) are supported"
        )

    content = await file.read()
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    if ws is None:
        raise HTTPException(status_code=400, detail="Excel file has no data")

    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    username_idx = None
    email_idx = None
    phone_idx = None
    role_name_idx = None

    for i, h in enumerate(headers):
        if h == "username":
            username_idx = i
        elif h == "email":
            email_idx = i
        elif h == "phone":
            phone_idx = i
        elif h in ["role_name", "role", "角色名"]:
            role_name_idx = i

    if username_idx is None:
        raise HTTPException(status_code=400, detail="Excel must have 'username' column")
    if role_name_idx is None:
        raise HTTPException(status_code=400, detail="Excel must have 'role_name' column")

    role_rows = db.query(Role.id, Role.name).all()
    role_name_map = {
        str(name).strip(): role_id
        for role_id, name in role_rows
        if name is not None and str(name).strip()
    }

    total = 0
    success = 0
    failed = 0
    errors = []

    for excel_row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue

        total += 1
        username = (
            str(row[username_idx]).strip() if row[username_idx] is not None else ""
        )
        email = (
            str(row[email_idx]).strip()
            if email_idx is not None and row[email_idx] is not None
            else ""
        )
        phone = (
            str(row[phone_idx]).strip()
            if phone_idx is not None and row[phone_idx] is not None
            else ""
        )
        role_name = (
            str(row[role_name_idx]).strip()
            if role_name_idx is not None and row[role_name_idx] is not None
            else ""
        )

        email = email if email else None
        phone = phone if phone else None

        if not username:
            failed += 1
            errors.append(f"Row {excel_row_num}: Username is required")
            continue

        if not role_name:
            failed += 1
            errors.append(f"Row {excel_row_num}: role_name is required")
            continue

        role_id = role_name_map.get(role_name)
        if not role_id:
            failed += 1
            errors.append(f"Row {excel_row_num}: Role '{role_name}' does not exist")
            continue

        existing = db.query(User).filter(User.username == username).first()
        if existing:
            failed += 1
            errors.append(f"Row {excel_row_num}: Username '{username}' already exists")
            continue

        if email:
            existing = db.query(User).filter(User.email == email).first()
            if existing:
                failed += 1
                errors.append(f"Row {excel_row_num}: Email '{email}' already exists")
                continue

        if phone:
            existing = db.query(User).filter(User.phone == phone).first()
            if existing:
                failed += 1
                errors.append(f"Row {excel_row_num}: Phone '{phone}' already exists")
                continue

        hashed_password = auth_service.get_password_hash(DEFAULT_PASSWORD)
        new_user = User(
            username=username,
            email=email,
            phone=phone,
            hashed_password=hashed_password,
        )
        db.add(new_user)
        db.flush()
        _replace_user_role(db, new_user.id, role_id)
        success += 1

    db.commit()
    return ImportResult(total=total, success=success, failed=failed, errors=errors)
